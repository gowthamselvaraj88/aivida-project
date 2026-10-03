import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # completely blank layout

    # Color Palette
    C_BRAND = RGBColor(0, 176, 116)         # Medical Emerald #00B074
    C_BRAND_DARK = RGBColor(0, 77, 51)       # Deep Emerald #004D33
    C_INK = RGBColor(17, 24, 39)             # Deep Slate #111827
    C_MUTED = RGBColor(107, 114, 128)        # Charcoal Gray #6B7280
    C_BG = RGBColor(244, 249, 246)           # Soft Mint Canvas #F4F9F6
    C_WHITE = RGBColor(255, 255, 255)
    C_CARD_BORDER = RGBColor(220, 235, 226)
    C_AMBER = RGBColor(245, 158, 11)
    C_BLUE = RGBColor(37, 99, 235)

    base_dir = os.path.dirname(os.path.abspath(__file__))
    shots_dir = os.path.join(base_dir, "screenshots")

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.color.rgb = C_BG
        return bg

    def add_header(slide, eyebrow, title, subtitle=None):
        # Header banner
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        # Eyebrow
        p0 = tf.paragraphs[0]
        p0.text = eyebrow.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = C_BRAND
        p0.space_after = Pt(2)
        
        # Title
        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = C_INK
        
        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.size = Pt(12)
            p2.font.color.rgb = C_MUTED
            p2.space_before = Pt(2)

    # ==========================================
    # SLIDE 1: Title Slide (Layman Welcome)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_BRAND_DARK
    bg1.line.color.rgb = C_BRAND_DARK

    # Accent decorative box
    dec = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    dec.fill.solid()
    dec.fill.fore_color.rgb = RGBColor(0, 92, 61)
    dec.line.color.rgb = RGBColor(0, 176, 116)
    dec.line.width = Pt(1.5)

    # Title text box
    tb = s1.shapes.add_textbox(Inches(1.4), Inches(1.4), Inches(10.5), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "AIVIDA TELEHEALTH PLATFORM"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(167, 243, 208)
    p.space_after = Pt(12)

    p = tf.add_paragraph()
    p.text = "Healthcare Simplified for Everyone"
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "A Non-Technical Guide: How Patients and Doctors Connect Seamlessly from Booking to Prescription."
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(209, 250, 229)
    p.space_after = Pt(30)

    p = tf.add_paragraph()
    p.text = "• Patient Care Portal  |  • Doctor Clinical Portal  |  • Happy Path Flow  |  • Interactive Block Diagram"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(254, 240, 138)

    # ==========================================
    # SLIDE 2: The Big Picture (What is AIvida in plain language?)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "Overview For Everyone", "The Big Picture: What is AIvida?", "Bridging the gap between home comfort and clinical expertise without delays or paperwork.")

    # 3 Layman Pillars Cards
    pillars = [
        ("No Waiting Rooms", "Patients connect directly with top-rated medical specialists from their phone, tablet, or PC in seconds.", C_BRAND),
        ("Clinical Precision", "Doctors have a full digital cockpit: live vital statistics, patient history, and 1-click video consultations.", C_BLUE),
        ("Zero Lost Paperwork", "Digital prescriptions, diagnosis summaries, and consultation records are stored securely and accessible forever.", C_AMBER)
    ]

    card_w = Inches(3.64)
    card_h = Inches(4.8)
    y_pos = Inches(1.8)

    for i, (ptitle, pdesc, pcol) in enumerate(pillars):
        x_pos = Inches(0.8 + i * 4.04)
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, y_pos, card_w, card_h)
        c.fill.solid()
        c.fill.fore_color.rgb = C_WHITE
        c.line.color.rgb = C_CARD_BORDER
        c.line.width = Pt(1.2)

        # Top color accent bar inside card
        bar = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos + Inches(0.3), y_pos + Inches(0.3), Inches(3.04), Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = pcol
        bar.line.color.rgb = pcol

        ctb = s2.shapes.add_textbox(x_pos + Inches(0.3), y_pos + Inches(0.6), Inches(3.04), Inches(3.8))
        ctf = ctb.text_frame
        ctf.word_wrap = True

        cp0 = ctf.paragraphs[0]
        cp0.text = f"PILLAR 0{i+1}"
        cp0.font.size = Pt(10)
        cp0.font.bold = True
        cp0.font.color.rgb = pcol
        cp0.space_after = Pt(8)

        cp1 = ctf.add_paragraph()
        cp1.text = ptitle
        cp1.font.size = Pt(20)
        cp1.font.bold = True
        cp1.font.color.rgb = C_INK
        cp1.space_after = Pt(12)

        cp2 = ctf.add_paragraph()
        cp2.text = pdesc
        cp2.font.size = Pt(13)
        cp2.font.color.rgb = C_MUTED
        cp2.space_after = Pt(16)

        cp3 = ctf.add_paragraph()
        cp3.text = "Key Layman Benefit:"
        cp3.font.size = Pt(11)
        cp3.font.bold = True
        cp3.font.color.rgb = C_INK
        cp3.space_after = Pt(4)

        benefits = [
            "Consult right from your living room without sitting in traffic.",
            "Physicians spend more time caring and less time typing.",
            "No fading paper slips; pharmacy can read electronic orders directly."
        ]
        cp4 = ctf.add_paragraph()
        cp4.text = benefits[i]
        cp4.font.size = Pt(11)
        cp4.font.color.rgb = C_BRAND_DARK

    # ==========================================
    # SLIDE 3: Patient Care Portal (Screenshots & Features)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "App 1: Patient Care Experience", "The Patient Care Portal (localhost:3001)", "Designed for utmost simplicity: find doctors, schedule appointments, and join calls with zero confusion.")

    # Left text box
    tb3 = s3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    tf3 = tb3.text_frame
    tf3.word_wrap = True

    p = tf3.paragraphs[0]
    p.text = "What Can a Patient Do Here?"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p.space_after = Pt(10)

    points = [
        ("Find the Right Doctor:", "Browse verified physicians with specialization, availability, and qualifications clearly displayed."),
        ("Pick Convenient Times:", "Select morning, afternoon, or evening slots with 1-click confirmation."),
        ("No Apps or Downloads:", "Video calls launch straight in the browser with crystal clear audio & video."),
        ("Health History & Prescriptions:", "View past consultations, active medication instructions, and upcoming checkups.")
    ]

    for title, desc in points:
        p = tf3.add_paragraph()
        p.text = f"• {title} "
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_BRAND
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_MUTED
        p.space_after = Pt(8)

    # Right Screenshots (Find Doctor & Home)
    p_img1 = os.path.join(shots_dir, "patient_doctors.png")
    if os.path.exists(p_img1):
        s3.shapes.add_picture(p_img1, Inches(5.6), Inches(1.8), width=Inches(7.0))

    # ==========================================
    # SLIDE 4: Doctor Clinical Portal (Screenshots & Features)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "App 2: Practice Workspace", "The Doctor Clinical Portal (localhost:3000)", "A powerful clinical command center that organizes the doctor's entire day into a single organized screen.")

    # Left text box
    tb4 = s4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    tf4 = tb4.text_frame
    tf4.word_wrap = True

    p = tf4.paragraphs[0]
    p.text = "What Can a Doctor Do Here?"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p.space_after = Pt(10)

    d_points = [
        ("Practice Availability Toggle:", "Go online/offline instantly with a single button to accept or pause walk-in telehealth slots."),
        ("At-a-Glance Telemetry:", "Review daily appointments, patient counts, vital statistics, and clinical telemetry cards."),
        ("Live Patient Queue:", "Organized consultation schedule with direct 'Start Call' buttons for scheduled time slots."),
        ("Digital Prescription Pad:", "Generate verified medical prescriptions with dosage, frequency, and instructions in seconds.")
    ]

    for title, desc in d_points:
        p = tf4.add_paragraph()
        p.text = f"• {title} "
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_BRAND
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_MUTED
        p.space_after = Pt(8)

    # Right Screenshot (Doctor Dashboard)
    d_img1 = os.path.join(shots_dir, "doctor_dashboard.png")
    if os.path.exists(d_img1):
        s4.shapes.add_picture(d_img1, Inches(5.6), Inches(1.8), width=Inches(7.0))

    # ==========================================
    # SLIDE 5: The Happy Path Flow (Step-by-Step Patient Journey)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "End-to-End User Experience", "The Happy Path: From Sickness to Prescription", "A seamless 5-step journey that takes less than 15 minutes from start to finish.")

    steps = [
        ("Step 1", "Search & Select", "Patient opens portal, searches doctor by specialty, and reviews credentials.", C_BLUE),
        ("Step 2", "Book & Confirm", "Selects an available slot. Database instantly reserves slot and sends confirmation.", C_BRAND),
        ("Step 3", "Join Call", "Both patient and doctor tap 'Join Consultation'. Video call launches instantly in HD.", C_AMBER),
        ("Step 4", "Live Telemetry", "Doctor reviews live health telemetry, discusses symptoms, and gives medical counsel.", C_BLUE),
        ("Step 5", "Rx Delivered", "Doctor signs digital prescription; patient receives medicine plan instantly on screen.", C_BRAND)
    ]

    step_w = Inches(2.26)
    step_h = Inches(4.8)
    for i, (snum, stitle, sdesc, scol) in enumerate(steps):
        sx = Inches(0.8 + i * 2.42)
        sy = Inches(1.8)

        scard = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, sy, step_w, step_h)
        scard.fill.solid()
        scard.fill.fore_color.rgb = C_WHITE
        scard.line.color.rgb = C_CARD_BORDER
        scard.line.width = Pt(1.2)

        # Step badge
        sbadge = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx + Inches(0.2), sy + Inches(0.25), Inches(1.0), Inches(0.35))
        sbadge.fill.solid()
        sbadge.fill.fore_color.rgb = scol
        sbadge.line.color.rgb = scol
        sbtf = sbadge.text_frame
        sbp = sbtf.paragraphs[0]
        sbp.text = snum
        sbp.font.size = Pt(11)
        sbp.font.bold = True
        sbp.font.color.rgb = C_WHITE
        sbp.alignment = PP_ALIGN.CENTER

        stb = s5.shapes.add_textbox(sx + Inches(0.2), sy + Inches(0.75), Inches(1.86), Inches(3.8))
        stf = stb.text_frame
        stf.word_wrap = True

        p = stf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_INK
        p.space_after = Pt(10)

        p = stf.add_paragraph()
        p.text = sdesc
        p.font.size = Pt(11)
        p.font.color.rgb = C_MUTED
        p.space_after = Pt(14)

        p = stf.add_paragraph()
        p.text = "Involved App:"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_INK
        p.space_after = Pt(2)

        apps = ["Patient Portal", "Cloud Server", "Both Portals", "Doctor Portal", "Both Portals"]
        p = stf.add_paragraph()
        p.text = apps[i]
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = scol

    # ==========================================
    # SLIDE 6: Working Block Diagram (Visual & Architectural Flow)
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "Under The Hood (Made Simple)", "Working Block Diagram: How Information Flows", "No complex engineering jargon — see how data, video, and medical records travel safely.")

    # 4 Main Blocks
    blocks = [
        ("PATIENT PORTAL\n(localhost:3001)", "• Patient Web App\n• Doctor Directory\n• Booking Engine\n• Prescription Viewer", Inches(0.8), C_BLUE),
        ("AIVIDA CORE SERVER\n(localhost:8080)", "• Go Telehealth API\n• Slot Management\n• Database (SQLite/SQL)\n• Auth & Patient Records", Inches(3.85), C_BRAND_DARK),
        ("WEBRTC GATEWAY\n(Live Video Engine)", "• Peer-to-Peer Video\n• Encrypted Audio Stream\n• Low-Latency Signaling\n• In-Call Chat", Inches(6.9), C_AMBER),
        ("DOCTOR PORTAL\n(localhost:3000)", "• Clinical Command Cockpit\n• Real-Time Telemetry\n• Digital Rx Generator\n• Practice Online Toggle", Inches(9.95), C_BRAND)
    ]

    bw = Inches(2.58)
    bh = Inches(3.2)
    by = Inches(1.8)

    for btitle, bbody, bx, bcol in blocks:
        blk = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, bw, bh)
        blk.fill.solid()
        blk.fill.fore_color.rgb = C_WHITE
        blk.line.color.rgb = bcol
        blk.line.width = Pt(2)

        # Header bar of block
        hbar = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, bw, Inches(0.75))
        hbar.fill.solid()
        hbar.fill.fore_color.rgb = bcol
        hbar.line.color.rgb = bcol
        htf = hbar.text_frame
        hp = htf.paragraphs[0]
        hp.text = btitle
        hp.font.size = Pt(11)
        hp.font.bold = True
        hp.font.color.rgb = C_WHITE
        hp.alignment = PP_ALIGN.CENTER

        # Body of block
        btb = s6.shapes.add_textbox(bx + Inches(0.2), by + Inches(0.85), bw - Inches(0.4), Inches(2.2))
        btf = btb.text_frame
        btf.word_wrap = True
        for line in bbody.split("\n"):
            p = btf.add_paragraph() if btf.paragraphs[0].text else btf.paragraphs[0]
            p.text = line
            p.font.size = Pt(11)
            p.font.color.rgb = C_INK if "•" not in line else C_MUTED
            p.space_after = Pt(4)

    # Bottom flow explanation card
    flow_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.25), Inches(11.73), Inches(1.6))
    flow_card.fill.solid()
    flow_card.fill.fore_color.rgb = C_WHITE
    flow_card.line.color.rgb = C_CARD_BORDER
    flow_card.line.width = Pt(1)

    ftb = s6.shapes.add_textbox(Inches(1.1), Inches(5.35), Inches(11.1), Inches(1.4))
    ftf = ftb.text_frame
    ftf.word_wrap = True

    p = ftf.paragraphs[0]
    p.text = "How The 3 Flows Operate In Real-Time:"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_BRAND
    p.space_after = Pt(4)

    flows = [
        ("1. Booking Flow: ", "Patient selects slot -> Core Server confirms -> Doctor dashboard schedule updates instantly."),
        ("2. Video Flow: ", "Signaling negotiates encrypted room -> WebRTC streams direct patient-doctor video with zero lag."),
        ("3. Prescription Flow: ", "Doctor clicks 'Issue Rx' -> Saved securely to database -> Available on Patient Portal immediately.")
    ]
    for fl_title, fl_text in flows:
        p = ftf.add_paragraph()
        p.text = fl_title
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = C_INK
        r = p.add_run()
        r.text = fl_text
        r.font.bold = False
        r.font.color.rgb = C_MUTED
        p.space_after = Pt(2)

    # ==========================================
    # SLIDE 7: Happy Path Screenshots Showcase (Prescriptions & Schedule)
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "Visual Evidence", "Inside the Live System: Appointments & Prescriptions", "Live captures from the running AIvida portals showing clean typography and medical hierarchy.")

    p_img2 = os.path.join(shots_dir, "doctor_prescriptions.png")
    p_img3 = os.path.join(shots_dir, "doctor_appointments.png")

    if os.path.exists(p_img2):
        s7.shapes.add_picture(p_img2, Inches(0.8), Inches(1.8), width=Inches(5.7))
    if os.path.exists(p_img3):
        s7.shapes.add_picture(p_img3, Inches(6.8), Inches(1.8), width=Inches(5.7))

    # Caption box
    ctb = s7.shapes.add_textbox(Inches(0.8), Inches(6.0), Inches(11.73), Inches(0.9))
    ctf = ctb.text_frame
    ctf.word_wrap = True
    p = ctf.paragraphs[0]
    p.text = "Left: Digital Prescriptions with distinct medication rows, dosage tags, and clear instructions."
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_INK
    p = ctf.add_paragraph()
    p.text = "Right: Consultation Schedule with Practice Availability toggle, time slots, and patient status badges."
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_INK

    # ==========================================
    # SLIDE 8: Conclusion & Layman Benefits
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_header(s8, "Summary & Impact", "Why AIvida Transforms Telehealth", "Key advantages for patients, healthcare professionals, and medical practices.")

    sum_cols = [
        ("For Patients", [
            "Access certified care in minutes",
            "Zero commuting, zero waiting rooms",
            "Automatic digital prescriptions",
            "Total privacy & medical data safety"
        ], C_BLUE),
        ("For Doctors", [
            "One-stop clinical command center",
            "Flexible practice online/offline control",
            "Instant electronic health records",
            "Reduced burnout & automated scheduling"
        ], C_BRAND),
        ("For Practices & Clinics", [
            "100% paperless workflow",
            "Higher patient retention & satisfaction",
            "Standardized clinical telemetry",
            "Scalable microservice architecture"
        ], C_AMBER)
    ]

    for i, (stitle, spoints, scol) in enumerate(sum_cols):
        cx = Inches(0.8 + i * 4.04)
        cy = Inches(1.8)
        scard = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, Inches(3.64), Inches(4.8))
        scard.fill.solid()
        scard.fill.fore_color.rgb = C_WHITE
        scard.line.color.rgb = C_CARD_BORDER
        scard.line.width = Pt(1.2)

        # Header bar
        bar = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + Inches(0.3), cy + Inches(0.3), Inches(3.04), Inches(0.55))
        bar.fill.solid()
        bar.fill.fore_color.rgb = scol
        bar.line.color.rgb = scol
        btf = bar.text_frame
        bp = btf.paragraphs[0]
        bp.text = stitle
        bp.font.size = Pt(14)
        bp.font.bold = True
        bp.font.color.rgb = C_WHITE
        bp.alignment = PP_ALIGN.CENTER

        tb = s8.shapes.add_textbox(cx + Inches(0.3), cy + Inches(1.05), Inches(3.04), Inches(3.5))
        tf = tb.text_frame
        tf.word_wrap = True

        for pt in spoints:
            p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
            p.text = f"✓  {pt}"
            p.font.size = Pt(13)
            p.font.color.rgb = C_INK
            p.space_after = Pt(14)

    output_path = os.path.join(base_dir, "AIvida_Telehealth_Presentation.pptx")
    prs.save(output_path)
    print(f"Successfully generated PowerPoint presentation at: {output_path}")

if __name__ == "__main__":
    create_deck()
